"""POST /api/v2/batch - price up to 100 instruments from JSON, CSV or XLSX."""
from __future__ import annotations

import csv
import io
import json
import time
from typing import Any, Optional

from fastapi import APIRouter, Request

from . import audit
from .common import (audit_block, error, internal_error, logger, metrics, price_cached, read_json, validate)
from .schemas import AdvancedRequest

router = APIRouter()

MAX_INSTRUMENTS = 100
MAX_UPLOAD_BYTES = 2 * 1024 * 1024
COLUMNS = ("face", "coupon", "yield", "years", "freq", "kind")
NUMERIC = ("face", "coupon", "yield", "years")
KIND_ALIASES = {"bond": "bond", "sukuk": "sukuk", "ijara": "sukuk", "ijara_sukuk": "sukuk",
                "conventional_bond": "bond"}
HEADER_ALIASES = {"profit_rate": "coupon", "profit": "coupon", "rental": "coupon", "yield_rate": "yield",
                  "type": "kind", "structure": "kind", "frequency": "freq", "tenor": "years"}


class BatchInputError(Exception):
    """Raised for unreadable or oversized batch input; message is user-facing."""


def _norm_header(name: Any) -> str:
    key = str(name or "").strip().lower().replace(" ", "_")
    return HEADER_ALIASES.get(key, key)


def _num(value: Any) -> float:
    """Convert CSV/spreadsheet text to a number (percent signs are rejected, rates are decimals)."""
    text = str(value).strip().replace(",", "")
    if text.endswith("%"):
        raise ValueError("percent")
    return float(text)


def _rows_from_table(header: list[Any], body_rows: list[list[Any]]) -> list[dict]:
    names = [_norm_header(h) for h in header]
    missing = [c for c in ("face", "coupon", "yield", "years") if c not in names]
    if missing:
        raise BatchInputError(
            f"Missing required column(s): {', '.join(missing)}. Expected columns: {', '.join(COLUMNS)}")
    rows = []
    for values in body_rows:
        if all(v is None or str(v).strip() == "" for v in values):
            continue
        row = {n: v for n, v in zip(names, values) if n in COLUMNS}
        rows.append({k: v for k, v in row.items() if v is not None and str(v).strip() != ""})
        if len(rows) > MAX_INSTRUMENTS:
            break
    return rows


def parse_csv(data: bytes) -> list[dict]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise BatchInputError("CSV must be UTF-8 encoded text") from exc
    table = list(csv.reader(io.StringIO(text)))
    if not table:
        raise BatchInputError("The CSV file is empty")
    return _rows_from_table(table[0], table[1:])


def parse_xlsx(data: bytes) -> list[dict]:
    try:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = wb.active
        table = [list(r) for r in sheet.iter_rows(min_row=1, max_row=MAX_INSTRUMENTS + 50, values_only=True)]
        wb.close()
    except Exception as exc:
        raise BatchInputError("The Excel file could not be read; upload a valid .xlsx workbook") from exc
    if not table:
        raise BatchInputError("The Excel file is empty")
    return _rows_from_table(table[0], table[1:])


def _coerce(data: dict) -> dict:
    """Turn CSV/XLSX text cells into numbers; JSON numbers pass through untouched."""
    for key in NUMERIC:
        if key in data and isinstance(data[key], str):
            data[key] = _num(data[key])
    if "freq" in data:
        freq = data["freq"]
        if isinstance(freq, str):
            freq = _num(freq)
        if isinstance(freq, float) and freq.is_integer():
            freq = int(freq)
        data["freq"] = freq
    for key in ("face", "years"):  # spreadsheets often hand back ints; keep strict floats happy
        if isinstance(data.get(key), int) and not isinstance(data[key], bool):
            data[key] = float(data[key])
    return data


def _price_row(index: int, row: Any) -> tuple[Optional[dict], Optional[dict]]:
    """Return (priced row, error row) for a single instrument."""
    if not isinstance(row, dict):
        return None, {"index": index, "error": "Each instrument must be an object"}
    data = dict(row)
    kind_raw = str(data.pop("kind", "bond") or "bond").strip().lower()
    kind = KIND_ALIASES.get(kind_raw)
    if kind is None:
        return None, {"index": index, "error": f"Unknown kind '{kind_raw}'; use bond or sukuk"}
    if "profit_rate" in data and "coupon" not in data:
        data["coupon"] = data.pop("profit_rate")
    try:
        data = _coerce(data)
    except ValueError:
        return None, {"index": index, "error": "Rates must be decimals (0.06 for 6%) and numeric"}
    data["structure"] = kind
    req, err = validate(AdvancedRequest, data)
    if err is not None:
        return None, {"index": index, "error": json.loads(err.body)["message"]}
    result, hit = price_cached(req)
    return {"index": index, "kind": kind, **metrics(result), "face": req.face, "yield": req.yield_rate,
            "cache_hit": hit}, None


def _wavg(rows: list[dict], key: str, weight: str) -> Optional[float]:
    total = sum(r[weight] for r in rows)
    return round(sum(r[key] * r[weight] for r in rows) / total, 8) if total else None


def price_batch(rows: list[Any]) -> dict:
    started = time.perf_counter()
    priced: list[dict] = []
    errors: list[dict] = []
    for index, row in enumerate(rows):
        ok, bad = _price_row(index, row)
        if ok:
            priced.append(ok)
        else:
            errors.append(bad)
    summary = {
        "count": len(rows),
        "priced": len(priced),
        "failed": len(errors),
        "total_face": round(sum(r["face"] for r in priced), 8),
        "total_market_value": round(sum(r["price"] for r in priced), 8),
        "weighted_avg_yield_by_face": _wavg(priced, "yield", "face"),
        "weighted_avg_yield_by_market_value": _wavg(priced, "yield", "price"),
        "weighted_avg_duration": _wavg(priced, "modified_duration", "price"),
        "total_dv01": round(sum(r["dv01"] for r in priced), 8),
        "cache_hits": sum(1 for r in priced if r["cache_hit"]),
    }
    return {"summary": summary, "results": priced, "errors": errors,
            "profile": {"v2_time_ms": round((time.perf_counter() - started) * 1000.0, 4)}}


async def _load_rows(request: Request) -> tuple[Optional[list], Optional[Any]]:
    ctype = (request.headers.get("content-type") or "").lower()
    try:
        if ctype.startswith("multipart/form-data"):
            form = await request.form()
            upload = form.get("file")
            if upload is None or not hasattr(upload, "read"):
                return None, error(400, ["Attach a .csv or .xlsx file in the 'file' field"])
            data = await upload.read()
            if len(data) > MAX_UPLOAD_BYTES:
                return None, error(400, ["Upload exceeds the 2 MB limit"])
            name = (upload.filename or "").lower()
            if name.endswith(".xlsx"):
                return parse_xlsx(data), None
            if name.endswith(".csv"):
                return parse_csv(data), None
            return None, error(400, ["Unsupported file type; upload .csv or .xlsx"])
        if ctype.startswith("text/csv"):
            data = await request.body()
            if len(data) > MAX_UPLOAD_BYTES:
                return None, error(400, ["Upload exceeds the 2 MB limit"])
            return parse_csv(data), None
        body, err = await read_json(request)
        if err:
            return None, err
        if isinstance(body, dict) and isinstance(body.get("instruments"), list):
            body = body["instruments"]
        if not isinstance(body, list):
            return None, error(400, ["Send a JSON array of instruments (or upload a CSV/XLSX file)"])
        return body, None
    except BatchInputError as exc:
        return None, error(400, [str(exc)])


@router.post("/batch")
async def batch(request: Request) -> Any:
    from .security import authenticate

    _, err = authenticate(request)  # optional JWT: anonymous allowed, invalid token rejected
    if err:
        return err
    rows, err = await _load_rows(request)
    if err:
        return err
    if not rows:
        return error(400, ["No instruments found; provide at least one row"])
    if len(rows) > MAX_INSTRUMENTS:
        return error(400, [f"Batch size exceeds the maximum of {MAX_INSTRUMENTS} instruments"])
    try:
        out = price_batch(rows)
        rec = audit.record("batch", {"instruments": rows}, out["summary"]["total_market_value"],
                           {"path": "python"}, out["summary"])
    except Exception:
        logger.exception("batch failed")
        return internal_error()
    return {**out, "audit": audit_block(rec)}
